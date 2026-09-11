"""SQLite persistence for uploaded video history."""

import os
from pathlib import Path
import sqlite3
import tempfile
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font


class UploadHistoryRepository:
    """Store upload records in a local SQLite database."""

    def __init__(self, database_path: Path) -> None:
        self.database_path = database_path
        self.database_path.parent.mkdir(parents=True, exist_ok=True)
        self._create_schema()

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.database_path)
        connection.row_factory = sqlite3.Row
        return connection

    def _create_schema(self) -> None:
        with self._connect() as connection:
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS uploads (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    youtube_video_id TEXT UNIQUE,
                    title TEXT NOT NULL,
                    file_path TEXT,
                    file_hash TEXT,
                    file_size INTEGER,
                    category TEXT,
                    playlist_id TEXT,
                    uploaded_at TEXT,
                    status TEXT NOT NULL,
                    source TEXT NOT NULL
                )
                """
            )
            connection.execute(
                "CREATE INDEX IF NOT EXISTS idx_uploads_file_hash "
                "ON uploads(file_hash)"
            )

    def upsert(self, record: dict[str, Any]) -> None:
        columns = (
            "youtube_video_id",
            "title",
            "file_path",
            "file_hash",
            "file_size",
            "category",
            "playlist_id",
            "uploaded_at",
            "status",
            "source",
        )
        values = [record.get(column) for column in columns]
        assignments = ", ".join(
            f"{column}=excluded.{column}"
            for column in columns
            if column != "youtube_video_id"
        )
        with self._connect() as connection:
            connection.execute(
                f"""
                INSERT INTO uploads ({', '.join(columns)})
                VALUES ({', '.join('?' for _ in columns)})
                ON CONFLICT(youtube_video_id) DO UPDATE SET {assignments}
                """,
                values,
            )

    def count(self) -> int:
        with self._connect() as connection:
            row = connection.execute("SELECT COUNT(*) AS count FROM uploads").fetchone()
        return int(row["count"])

    def find_by_video_id(self, video_id: str) -> sqlite3.Row | None:
        with self._connect() as connection:
            return connection.execute(
                "SELECT * FROM uploads WHERE youtube_video_id = ?",
                (video_id,),
            ).fetchone()

    def export_to_excel(self, output_path: Path) -> int:
        """Export every upload history record to an Excel workbook."""
        database_columns = (
            "id",
            "youtube_video_id",
            "title",
            "file_path",
            "file_hash",
            "file_size",
            "category",
            "playlist_id",
            "uploaded_at",
            "status",
            "source",
        )
        title_index = database_columns.index("title") + 1
        columns = (
            *database_columns[:title_index],
            "video_url",
            *database_columns[title_index:],
        )
        with self._connect() as connection:
            rows = connection.execute(
                f"SELECT {', '.join(database_columns)} FROM uploads "
                "ORDER BY uploaded_at DESC, id DESC"
            ).fetchall()

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Upload History"
        worksheet.append(columns)
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            video_url = (
                f"https://www.youtube.com/watch?v={row['youtube_video_id']}"
                if row["youtube_video_id"]
                else None
            )
            values = [row[column] for column in database_columns]
            values.insert(title_index, video_url)
            worksheet.append(values)

        video_url_column = title_index + 1
        for cell in next(
            worksheet.iter_cols(
                min_col=video_url_column,
                max_col=video_url_column,
                min_row=2,
            ),
            (),
        ):
            if cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for column_cells in worksheet.columns:
            column_letter = column_cells[0].column_letter
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            worksheet.column_dimensions[column_letter].width = min(max_length + 2, 60)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        temporary_path: Path | None = None
        try:
            with tempfile.NamedTemporaryFile(
                dir=output_path.parent,
                prefix=f".{output_path.stem}-",
                suffix=output_path.suffix,
                delete=False,
            ) as temporary_file:
                temporary_path = Path(temporary_file.name)
            workbook.save(temporary_path)
            os.replace(temporary_path, output_path)
        finally:
            if temporary_path and temporary_path.exists():
                temporary_path.unlink()

        return len(rows)

    def report(self, recent_limit: int = 10) -> dict[str, Any]:
        """Return aggregate counts and the most recent history records."""
        with self._connect() as connection:
            total = connection.execute(
                "SELECT COUNT(*) AS count FROM uploads"
            ).fetchone()["count"]
            grouped_queries = {
                "status": "SELECT status AS name, COUNT(*) AS count "
                "FROM uploads GROUP BY status ORDER BY status",
                "category": "SELECT COALESCE(category, '') AS name, COUNT(*) AS count "
                "FROM uploads GROUP BY category ORDER BY category",
                "source": "SELECT source AS name, COUNT(*) AS count "
                "FROM uploads GROUP BY source ORDER BY source",
            }
            groups = {
                group_name: [dict(row) for row in connection.execute(query)]
                for group_name, query in grouped_queries.items()
            }
            recent = [
                dict(row)
                for row in connection.execute(
                    "SELECT title, youtube_video_id, category, status, uploaded_at "
                    "FROM uploads ORDER BY uploaded_at DESC, id DESC LIMIT ?",
                    (recent_limit,),
                )
            ]

        return {"total": int(total), **groups, "recent": recent}

    def find_successful_by_hash(self, file_hash: str) -> sqlite3.Row | None:
        with self._connect() as connection:
            return connection.execute(
                "SELECT * FROM uploads WHERE file_hash = ? AND status = 'uploaded' "
                "LIMIT 1",
                (file_hash,),
            ).fetchone()

    def find_successful_by_title(self, title: str) -> sqlite3.Row | None:
        with self._connect() as connection:
            return connection.execute(
                "SELECT * FROM uploads WHERE title = ? AND status = 'uploaded' "
                "ORDER BY uploaded_at DESC, id DESC LIMIT 1",
                (title,),
            ).fetchone()

    def uploaded_video_ids(self) -> set[str]:
        with self._connect() as connection:
            rows = connection.execute(
                "SELECT youtube_video_id FROM uploads "
                "WHERE status = 'uploaded' AND youtube_video_id IS NOT NULL"
            )
        return {row[0] for row in rows}

    def remove_uploaded_ids(self, video_ids: set[str]) -> int:
        if not video_ids:
            return 0
        with self._connect() as connection:
            placeholders = ", ".join("?" for _ in video_ids)
            cursor = connection.execute(
                "DELETE FROM uploads WHERE status = 'uploaded' "
                f"AND youtube_video_id IN ({placeholders})",
                tuple(video_ids),
            )
        return cursor.rowcount